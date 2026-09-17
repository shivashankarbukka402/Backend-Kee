"""
Generic Record Excel Exporter

Exports any dataclass records (item prices, opening stock, image mappings) into
a single ``.xlsx`` file that matches an ERPNext import template.

This is the reusable exporter infrastructure for the Phase 4 catalog-enrichment
records and any future generator record. Unlike the item-specific
:class:`~master_data.exporters.excel_exporter.ExcelExporter`, the column template
and record type are fully driven by configuration rather than being hard-coded
to the :class:`~master_data.item_models.Item` model, so the same writer serves
every record shape.

It performs no ERPNext import and never touches the ERPNext database.
"""

from __future__ import annotations

import logging
from dataclasses import fields, is_dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from ..config import EnrichmentExportConfig, ExportConfig
from ..logging_setup import get_logger
from .excel_exporter import ExportReport, ValidationError


class RecordExcelExporter:
    """
    Writes a single ERPNext import ``.xlsx`` file from ordered dataclass records.
    """

    def __init__(
        self,
        config: ExportConfig,
        logger: logging.Logger | None = None,
    ) -> None:
        self._config = config
        self._logger = logger or get_logger(self.__class__.__name__)

    def export(
        self,
        records: list[object] | tuple[object, ...],
        spec: EnrichmentExportConfig,
        *,
        output_dir: Path | None = None,
    ) -> ExportReport:
        """
        Write the given records into one workbook and return the export report.

        Parameters
        ----------
        records:
            The dataclass records to export.
        spec:
            Centralized column/required-fields/target-file rules for the record
            type being exported.
        output_dir:
            Base output directory (defaults to the configured export directory).
        """
        target_dir = (output_dir or self._config.output_dir) / spec.subdirectory
        target_dir.mkdir(parents=True, exist_ok=True)

        total = len(records)
        valid, errors = self._validated_records(records, spec)
        exported = len(valid)
        path = self._write_workbook(valid, spec, target_dir)
        self._logger.info("Exported %d record(s) to %s", exported, path)

        return ExportReport(
            total_generated=total,
            exported=exported,
            skipped=total - exported,
            files_generated=(path,),
            errors=tuple(errors),
        )

    # ------------------------------------------------------------------ #
    # Validation
    # ------------------------------------------------------------------ #

    def _validated_records(
        self,
        records: list[object] | tuple[object, ...],
        spec: EnrichmentExportConfig,
    ) -> tuple[list[object], list[ValidationError]]:
        """
        Keep only the records whose required attributes are all non-blank.
        """
        valid: list[object] = []
        errors: list[ValidationError] = []

        for record in records:
            code = self._record_code(record, spec)
            missing = self._missing_required(record, spec)
            if missing:
                reasons = ", ".join(missing)
                errors.append(ValidationError(item_code=code, reason=reasons))
                self._logger.warning(
                    "Skipping %s: missing required field(s) %s", code, reasons
                )
                continue
            valid.append(record)
        return valid, errors

    def _missing_required(
        self,
        record: object,
        spec: EnrichmentExportConfig,
    ) -> list[str]:
        """
        Return the required attribute names that are blank on the record.
        """
        return [
            attribute
            for attribute in spec.required_attributes
            if not str(getattr(record, attribute, "")).strip()
        ]

    def _record_code(self, record: object, spec: EnrichmentExportConfig) -> str:
        """Best-effort human identifier for a record (its Item Code if present)."""
        if "item_code" in spec.required_attributes:
            return str(getattr(record, "item_code", ""))
        return ""

    # ------------------------------------------------------------------ #
    # Workbook generation
    # ------------------------------------------------------------------ #

    def _write_workbook(
        self,
        records: list[object],
        spec: EnrichmentExportConfig,
        target_dir: Path,
    ) -> Path:
        """
        Write one workbook and return its path.
        """
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = spec.sheet_title

        self._write_header(sheet, spec)
        self._write_rows(sheet, records, spec)

        if self._config.auto_size_columns:
            self._auto_size(sheet, spec)

        if self._config.freeze_header_row:
            sheet.freeze_panes = "A2"

        path = target_dir / spec.filename
        workbook.save(path)
        workbook.close()
        return path

    def _write_header(self, sheet, spec: EnrichmentExportConfig) -> None:
        """Write the column headers into the first row."""
        for column_index, column in enumerate(spec.columns, start=1):
            cell = sheet.cell(row=1, column=column_index, value=column.name)
            if self._config.bold_header:
                cell.font = Font(bold=True)

    def _write_rows(
        self,
        sheet,
        records: list[object],
        spec: EnrichmentExportConfig,
    ) -> None:
        """Write each record as one data row below the header."""
        for row_index, record in enumerate(records, start=2):
            for column_index, column in enumerate(spec.columns, start=1):
                sheet.cell(
                    row=row_index,
                    column=column_index,
                    value=self._value(record, column.attribute),
                )

    def _value(self, record: object, attribute: str) -> object:
        """Serialize a record attribute for the spreadsheet."""
        if is_dataclass(record):
            allowed = {field.name for field in fields(record)}
        else:
            allowed = set()
        value = getattr(record, attribute, None) if attribute in allowed else None
        if isinstance(value, bool):
            return "1" if value else "0"
        if isinstance(value, (int, float)):
            return value
        return str(value) if value is not None else ""

    def _auto_size(self, sheet, spec: EnrichmentExportConfig) -> None:
        """Resize each column to fit its widest cell, with a sensible maximum."""
        for column_index, column in enumerate(spec.columns, start=1):
            widths = [len(column.name)]
            letter = get_column_letter(column_index)
            for cell in sheet[letter][1:]:
                if cell.value is not None:
                    widths.append(len(str(cell.value)))
            width = min(max(widths) + 2, 80)
            sheet.column_dimensions[letter].width = width


def build_record_exporter(
    config: ExportConfig,
    logger: logging.Logger | None = None,
) -> RecordExcelExporter:
    """
    Build a default generic record exporter from an export configuration.
    """
    return RecordExcelExporter(config=config, logger=logger)
